# -*- coding: utf-8 -*-
"""Builds the Shared National Talent package (proposal, NDA, internal contract draft) as PDFs.
Usage: python3 gen_nst.py [WORKDIR]   (WORKDIR holds fonts, recalculated workbook, html; default ./_work)
Prices are read from Shared-National-Talent-Pricing-INTERNAL.xlsx (recalculated by LibreOffice), never typed here."""
import os, sys, re, html, base64, subprocess, shutil, urllib.request
import openpyxl
import nst_content as C

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(HERE, '_work')
OUT = os.path.join(HERE, '..', 'national-shared-talent-out')
os.makedirs(WORK, exist_ok=True); os.makedirs(OUT, exist_ok=True)
CHROME = next(p for p in ('/opt/pw-browsers/chromium-1194/chrome-linux/chrome', '/opt/pw-browsers/chromium') if os.path.exists(p))
FONTS = {'ar400': 'ibmplexsansarabic/v15/Qw3CZRtWPQCuHme67tEYUIx3Kh0PHR9N6bs6.ttf',
         'ar600': 'ibmplexsansarabic/v15/Qw3NZRtWPQCuHme67tEYUIx3Kh0PHR9N6YPi-NCT.ttf'}
NAVY, INK, MUTE, LINE, PALE = '#0B1B5A', '#1B2340', '#6A7596', '#DDE1EE', '#F7F8FD'


def font_css():
    css = ''
    for k, u in FONTS.items():
        f = os.path.join(WORK, k + '.ttf')
        if not os.path.exists(f):
            urllib.request.urlretrieve('https://fonts.gstatic.com/s/' + u, f)
        w = 400 if k.endswith('400') else 600
        css += "@font-face{font-family:'PlexAr';font-weight:%d;src:url(data:font/ttf;base64,%s)}" % (w, base64.b64encode(open(f, 'rb').read()).decode())
    return css


def logo_uri():
    from PIL import Image
    im = Image.open(os.path.join(HERE, '..', '..', '..', '..', 'site', 'assets', 'img', 'logo.png')); im.thumbnail((700, 200))
    f = os.path.join(WORK, 'logo.png'); im.save(f)
    return 'data:image/png;base64,' + base64.b64encode(open(f, 'rb').read()).decode()


def prices():
    src = os.path.join(HERE, 'Shared-National-Talent-Pricing-INTERNAL.xlsx')
    rc = os.path.join(WORK, 'recalc'); os.makedirs(rc, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', rc, src], check=True, capture_output=True)
    wb = openpyxl.load_workbook(os.path.join(rc, os.path.basename(src)), data_only=True)
    r, h = wb['Rate card'], wb['Hour packages']
    out = {}
    for i in range(9):
        name = r.cell(12 + i, 1).value
        out[name] = dict(rate=r.cell(12 + i, 13).value, p20=h.cell(4 + i, 3).value, p40=h.cell(4 + i, 4).value,
                         p80=h.cell(4 + i, 5).value, p120=h.cell(4 + i, 6).value, full=h.cell(4 + i, 7).value)
    return out


esc = lambda s: html.escape(s, quote=False)
PH = re.compile(r'\[\[(.+?)\]\]')


def ph(s):
    out = PH.sub(lambda m: '<mark class="ph">[%s]</mark>' % m.group(1), esc(s))
    if re.search('[\u0600-\u06FF]', s) and s.rstrip().endswith('.'):
        out += '\u200f'   # sentence-final full stop stays at the end of the Arabic line (sales standard)
    return out


