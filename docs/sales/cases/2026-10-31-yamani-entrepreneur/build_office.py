# -*- coding: utf-8 -*-
"""DOCX للعقد وNDA، وPPTX (شرائح) لعرض السعر — بالهوية نفسها.
python3 build_office.py <outdir> [--final]
"""
import os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt as PPt
from content import *

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(ROOT, "docs/sales/templates/assets/business-partner-logo-clean.png")
FINAL = "--final" in sys.argv
PHONE = os.environ.get("BP_PHONE", "0530540231") if FINAL else "[●]"
IBAN = os.environ.get("BP_IBAN", "[IBAN]") if FINAL else "[IBAN]"
BANK = os.environ.get("BP_BANK", "[●]") if FINAL else "[●]"

NAVY, NAVY2, ACC, ACCD, TXT, TXT2, BG, LINE = "0B1B5A", "13246E", "1E3A9E", "9DB4FF", "1B2340", "4A5578", "F7F8FD", "DDE1EE"
F_AR, F_EN, F_H, F_M = "IBM Plex Sans Arabic", "IBM Plex Sans Arabic", "Playfair Display", "IBM Plex Mono"

def rgb(h): return RGBColor.from_string(h)

# ------------------------------------------------------------------ أدوات DOCX
def set_fonts(run, name, size, bold=False, color=TXT, rtl=False):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if rtl:
        el = OxmlElement("w:rtl"); rpr.append(el)
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rpr.append(szcs)
        if bold:
            bcs = OxmlElement("w:bCs"); rpr.append(bcs)

def para(cell_or_doc, text="", size=10, bold=False, color=TXT, align=None, font=None, rtl=False, space_after=2, first=False, line=1.25):
    p = cell_or_doc.paragraphs[0] if (first and hasattr(cell_or_doc, "paragraphs") and cell_or_doc.paragraphs and not cell_or_doc.paragraphs[0].text) else cell_or_doc.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = line
    if rtl:
        ppr = p._p.get_or_add_pPr(); b = OxmlElement("w:bidi"); ppr.insert(0, b)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # في فقرة bidi: «left» = بداية السطر (اليمين)
    elif align is not None:
        p.alignment = align
    if text:
        add_runs(p, text, size, bold, color, font or (F_AR if rtl else F_EN), rtl)
    return p

def add_runs(p, text, size, bold, color, font, rtl):
    """يدعم **غامق** داخل النص."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part: continue
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]; b = True
        r = p.add_run(part); set_fonts(r, font, size, b, color, rtl)

def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcpr.append(sh)

def borders(cell, color=LINE, sz=4, sides=("top", "left", "bottom", "right")):
    tcpr = cell._tc.get_or_add_tcPr(); tb = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + s)
        if s in sides:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        tb.append(e)
    tcpr.append(tb)

def cell_margins(cell, t=60, b=60, l=110, r=110):
    tcpr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, v in (("top", t), ("left", l), ("bottom", b), ("right", r)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcpr.append(m)

def col_widths(table, widths_cm):
    table.autofit = False
    tblpr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))

def no_split(row):
    trpr = row._tr.get_or_add_trPr(); e = OxmlElement("w:cantSplit"); trpr.append(e)

def set_cell_bidi_table(table):
    pass

def doc_setup(title_en, title_ar):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(3.0), Cm(3.0), Cm(1.9), Cm(1.9)
    s.header_distance, s.footer_distance = Cm(0.9), Cm(0.8)
    st = d.styles["Normal"]; st.font.name = F_EN; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:cs"), F_AR)
    # الرأس: الشعار يسار + العنوان
    h = s.header; ht = h.add_table(1, 2, Cm(17.2)); col_widths(ht, [6.5, 10.7])
    c0, c1 = ht.rows[0].cells
    c0.paragraphs[0].add_run().add_picture(LOGO, width=Cm(5.4))
    para(c1, title_ar, 15, True, NAVY, rtl=True, first=True, space_after=0)
    para(c1, title_en, 10, False, TXT2, align=WD_ALIGN_PARAGRAPH.RIGHT, font=F_H, space_after=0)
    for c in (c0, c1): borders(c, NAVY, 18, ("bottom",))
    h.paragraphs[0]._p.getparent().remove(h.paragraphs[0]._p)
    # التذييل
    f = s.footer; ft = f.add_table(1, 3, Cm(17.2)); col_widths(ft, [6.0, 6.0, 5.2])
    a, c, w = ft.rows[0].cells
    for t in ("National Address: RDAA9214", "9214 King Khaled Branch Rd, Al Asemah", "Ad Diriyah 13714, Saudi Arabia"):
        para(a, t, 7.5, False, TXT2, first=(t.startswith("National")), space_after=0, line=1.1)
    for k, t in enumerate((PHONE, "business@businesspartner.sa", "www.businesspartner.sa")):
        para(c, t, 7.5, False, TXT2, align=WD_ALIGN_PARAGRAPH.CENTER, first=(k == 0), space_after=0, line=1.1)
    # رقم الصفحة
    p = para(c, "", 7.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    for instr in ("PAGE", None, "NUMPAGES"):
        r = p.add_run(); set_fonts(r, F_M, 7.5, False, "8A93B0")
        if instr is None:
            r.text = " / "; continue
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
        fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "end"); r._r.append(fc2)
    w.paragraphs[0].add_run().add_picture(LOGO, width=Cm(3.6)); w.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f.paragraphs[0]._p.getparent().remove(f.paragraphs[0]._p)
    return d

def cover_block(d, kind_en, kind_ar, t_en, t_ar, s_en, s_ar):
    t = d.add_table(1, 1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; col_widths(t, [17.2])
    c = t.rows[0].cells[0]; shade(c, NAVY); borders(c, NAVY, 4); cell_margins(c, 700, 700, 500, 500)
    para(c, kind_en.upper(), 9, False, ACCD, font=F_M, first=True, space_after=2)
    para(c, t_en, 26, True, "FFFFFF", font=F_H, space_after=6, line=1.1)
    para(c, s_en, 11, False, "EEF1FB", space_after=18)
    para(c, kind_ar, 11, False, ACCD, rtl=True, space_after=2)
    para(c, t_ar, 24, True, "FFFFFF", rtl=True, space_after=6, line=1.2)
    para(c, ar(s_ar), 12, False, "EEF1FB", rtl=True, space_after=30)
    para(c, "Parties / الأطراف:  Business Partner Solutions · %s · %s" % (CLIENT_EN, CLIENT_AR), 10, False, "FFFFFF", space_after=2)
    para(c, "Date / التاريخ:  %s · %s" % (DATE_EN, DATE_AR), 10, False, "FFFFFF", space_after=2)
    para(c, "Ref / المرجع:  BP-FI-05", 10, False, "FFFFFF", space_after=0)

def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def kv_table(d, head_en, head_ar, rows):
    t = d.add_table(len(rows) + 1, 3); col_widths(t, [5.2, 6.8, 5.2])
    hd = t.rows[0].cells; m = hd[0].merge(hd[2]); shade(m, NAVY); borders(m, NAVY)
    para(m, "%s   |   %s" % (head_en, head_ar), 10, True, "FFFFFF", first=True, space_after=0)
    for i, (le, v, la) in enumerate(rows, 1):
        cs = t.rows[i].cells
        for c in cs: borders(c); cell_margins(c)
        if i % 2 == 0:
            for c in cs: shade(c, BG)
        para(cs[0], le, 9, False, TXT2, first=True, space_after=0)
        para(cs[1], v, 9.5, False, TXT, align=WD_ALIGN_PARAGRAPH.CENTER, first=True, space_after=0)
        para(cs[2], la, 9, False, TXT2, rtl=True, first=True, space_after=0)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def clauses_table(d, rows, start=2):
    t = d.add_table(len(rows), 2); col_widths(t, [8.6, 8.6])
    for i, r in enumerate(rows):
        t_en, t_ar, b_en, b_ar = r; no = start + i
        en, a = t.rows[i].cells
        no_split(t.rows[i])
        for c in (en, a): borders(c); cell_margins(c, 80, 80, 130, 130)
        if no % 2 == 1:
            for c in (en, a): shade(c, BG)
        para(en, "%d. %s" % (no, t_en), 9.5, True, NAVY, first=True, space_after=2)
        para(en, b_en.replace("&amp;", "&"), 9, False, TXT, space_after=0, line=1.3)
        para(a, "%d. %s" % (no, t_ar), 9.5, True, NAVY, rtl=True, first=True, space_after=2)
        para(a, ar(b_ar), 9.5, False, TXT, rtl=True, space_after=0, line=1.3)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def sign_table(d):
    t = d.add_table(1, 2); col_widths(t, [8.6, 8.6])
    for c, (te, ta) in zip(t.rows[0].cells, (("First Party", "الطرف الأول"), ("Second Party", "الطرف الثاني"))):
        borders(c, NAVY, 8); cell_margins(c, 0, 80, 0, 0)
        inner = c.add_table(5, 3); col_widths(inner, [2.4, 3.8, 2.2])
        h = inner.rows[0].cells; m = h[0].merge(h[2]); shade(m, NAVY)
        para(m, "%s     %s" % (te, ta), 10, True, "FFFFFF", first=True, space_after=0)
        for i, (le, la) in enumerate((("Name", "الاسم"), ("Title", "الصفة"), ("Date", "التاريخ")), 1):
            cs = inner.rows[i].cells
            para(cs[0], le, 9, False, TXT2, first=True, space_after=0)
            para(cs[1], "[●]", 9.5, False, TXT, align=WD_ALIGN_PARAGRAPH.CENTER, first=True, space_after=0)
            para(cs[2], la, 9, False, TXT2, rtl=True, first=True, space_after=0)
            for x in cs: borders(x, LINE, 4, ("bottom",)); cell_margins(x)
        sg = inner.rows[4].cells; mm = sg[0].merge(sg[2])
        for _ in range(4): para(mm, "", 10, space_after=6)
        para(mm, "Signature & Stamp · التوقيع والختم", 8, False, "8A93B0", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        borders(mm, "8A93B0", 4)
        # الخلية الأولى الفارغة الافتراضية
        c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p) if c.paragraphs and not c.paragraphs[0].text else None
        c.add_paragraph().paragraph_format.space_after = Pt(0)

def closing_block(d):
    page_break(d)
    t = d.add_table(1, 1); col_widths(t, [17.2])
    c = t.rows[0].cells[0]; shade(c, NAVY); cell_margins(c, 700, 700, 500, 500)
    para(c, "PARTNERING FOR YOUR SUCCESS", 9, False, ACCD, font=F_M, first=True, space_after=4)
    para(c, "Business Partner Solutions", 24, True, "FFFFFF", font=F_H, space_after=2)
    para(c, "بزنس بارتنر سلوشنز", 14, False, "FFFFFF", rtl=True, space_after=14)
    for le, v, la in (("Unified No.", "7038825860", "الرقم الموحد"), ("VAT No.", "312079341500003", "الرقم الضريبي"),
                      ("National Address", "RDAA9214 · 9214 King Khaled Branch Rd, Al Asemah, Ad Diriyah 13714", "العنوان الوطني"),
                      ("Mobile", PHONE, "الجوال"), ("Email", "business@businesspartner.sa", "الإيميل"), ("Website", "www.businesspartner.sa", "الموقع")):
        para(c, "%s  ·  %s  ·  %s" % (le, v, la), 10, False, "FFFFFF", space_after=4)

def build_docx(kind, rows, t_en, t_ar, cover, out):
    d = doc_setup(t_en, t_ar)
    cover_block(d, *cover)
    page_break(d)
    para(d, "1. الأطراف / Parties", 11, True, NAVY, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4)
    kv_table(d, "First Party (Service Provider)", "الطرف الأول (مزوّد الخدمة)", [
        ("Name", "Business Partner Solutions · شركة بزنس بارتنر سلوشنز", "الاسم"),
        ("Unified / CR No.", "7038825860", "الرقم الموحد / السجل التجاري"),
        ("VAT No.", "312079341500003", "الرقم الضريبي"),
        ("National Address", "RDAA9214", "العنوان الوطني"),
        ("Represented by", "[●]", "يمثلها"), ("Title", "[●]", "الصفة")])
    kv_table(d, "Second Party (Client)", "الطرف الثاني (العميل)", [
        ("Name", "%s · %s — full name as in ID: [●]" % (CLIENT_EN, CLIENT_AR), "الاسم"),
        ("ID / Iqama / Passport No.", "[●]", "رقم الهوية / الإقامة / الجواز"),
        ("Nationality", "[●]", "الجنسية"),
        ("National / residence address", "[●]", "العنوان الوطني / السكن"),
        ("Capacity", "Founder of the company to be established · مؤسس الشركة المراد تأسيسها", "الصفة")])
    clauses_table(d, rows)
    sign_table(d)
    closing_block(d)
    d.save(out)

def cover_of(kind):
    if kind == "agreement":
        return ("Service Agreement", "عقد خدمة", C_TITLE_EN, C_TITLE_AR,
                "Registering the investment under the entrepreneurship category and establishing the company.",
                "تسجيل الاستثمار في فئة ريادة الأعمال وتأسيس الشركة.")
    return ("Confidentiality", "سرية", N_TITLE_EN, N_TITLE_AR,
            "A mutual agreement protecting the confidential information exchanged between the parties.",
            "اتفاقية متبادلة لحماية المعلومات السرية المتبادلة بين الطرفين.")

# ------------------------------------------------------------------ PPTX
SW, SH = Inches(13.333), Inches(7.5)

def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = PRGB.from_string(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = PRGB.from_string(line); s.line.width = PPt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def txt(sl, x, y, w, h, text, size=14, bold=False, color=TXT, font=F_EN, rtl=False, align=None, anchor=MSO_ANCHOR.TOP, line=1.15):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line
        parts = re.split(r"(\*\*.+?\*\*)", ln)
        for part in parts:
            if not part: continue
            b = bold
            if part.startswith("**"): part = part[2:-2]; b = True
            r = p.add_run(); r.text = part
            r.font.size = PPt(size); r.font.bold = b; r.font.name = font; r.font.color.rgb = PRGB.from_string(color)
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:cs",):
                e = rPr.makeelement(qn(tag), {"typeface": font}); rPr.append(e)
        pPr = p._p.get_or_add_pPr()
        if rtl:
            pPr.set("rtl", "1"); pPr.set("algn", "r")
        elif align == "c":
            p.alignment = PP_ALIGN.CENTER
        elif align == "r":
            p.alignment = PP_ALIGN.RIGHT
    return tb

def slide(prs, bg):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    box(sl, 0, 0, 13.333, 7.5, bg)
    return sl

def footer_strip(sl, i, n, dark=False):
    c = ACCD if dark else TXT2
    box(sl, 0.6, 6.88, 12.13, 0.02, ACCD if dark else NAVY)
    txt(sl, 0.6, 6.95, 6, 0.3, "National Address: RDAA9214 · 9214 King Khaled Branch Rd, Al Asemah, Ad Diriyah 13714", 8, False, c)
    txt(sl, 6.3, 6.95, 6.43, 0.3, "%s · business@businesspartner.sa · www.businesspartner.sa · %d / %d" % (PHONE, i, n), 8, False, c, align="r")

def title_bar(sl, en, a):
    txt(sl, 0.6, 0.45, 6.2, 0.9, en, 24, True, NAVY, font=F_H)
    txt(sl, 7.0, 0.45, 5.73, 0.9, a, 28, True, NAVY, rtl=True)
    box(sl, 0.6, 1.35, 12.13, 0.04, NAVY)

def build_pptx(out):
    prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
    N = 5
    # 1 الغلاف
    sl = slide(prs, NAVY)
    box(sl, 0, 0, 13.333, 1.55, "FDFDFC")
    sl.shapes.add_picture(LOGO, Inches(0.6), Inches(0.42), width=Inches(4.2))
    txt(sl, 0.6, 2.45, 5.9, 0.3, "QUOTATION", 12, False, ACCD, font=F_M)
    txt(sl, 0.6, 2.8, 5.9, 1.6, "Entrepreneurship-Licence Company Formation", 30, True, "FFFFFF", font=F_H, line=1.1)
    txt(sl, 0.6, 4.4, 5.9, 1.2, "Registering your investment under the entrepreneurship category with the Ministry of Investment, then establishing your company and registering it with the competent authorities.", 14, False, "EEF1FB")
    txt(sl, 6.9, 2.45, 5.83, 0.3, "عرض سعر", 14, False, ACCD, rtl=True)
    txt(sl, 6.9, 2.8, 5.83, 1.6, "تأسيس شركة برخصة ريادة الأعمال", 34, True, "FFFFFF", rtl=True, line=1.1)
    txt(sl, 6.9, 4.4, 5.83, 1.2, ar("تسجيل استثمارك في فئة ريادة الأعمال لدى وزارة الاستثمار، ثم تأسيس شركتك وتسجيلها لدى الجهات المختصة."), 16, False, "EEF1FB", rtl=True)
    for k, (le, v) in enumerate((("To / إلى السادة", "%s · %s" % (CLIENT_EN, CLIENT_AR)), ("Date / التاريخ", "%s · %s" % (DATE_EN, DATE_AR)), ("Validity / صلاحية العرض", "14 days · 14 يوماً"))):
        x = 0.6 + k * 4.1
        box(sl, x, 5.75, 3.9, 0.03, ACCD)
        txt(sl, x, 5.82, 3.9, 0.3, le, 10, False, ACCD)
        txt(sl, x, 6.1, 3.9, 0.5, v, 12, False, "FFFFFF")
    txt(sl, 0.6, 6.95, 12.1, 0.3, "BP-FI-05 · Foreign Investment · الاستثمار الأجنبي", 9, False, ACCD, font=F_M)

    # 2 النطاق
    sl = slide(prs, "FFFFFF"); title_bar(sl, "Scope — included in the fee", "النطاق — شامل ضمن الأتعاب")
    for k, (te, ta, de, da) in enumerate(SCOPE):
        y = 1.6 + k * 0.74
        box(sl, 0.6, y, 12.13, 0.66, BG if k % 2 == 0 else "FFFFFF", LINE)
        txt(sl, 0.7, y + 0.07, 0.4, 0.5, str(k + 1), 16, True, ACC, font=F_M)
        txt(sl, 1.15, y + 0.04, 5.3, 0.6, "**%s** — %s" % (te, de), 10, False, TXT, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
        txt(sl, 6.65, y + 0.04, 5.6, 0.6, "**%s** — %s" % (ta, ar(da)), 11, False, TXT, rtl=True, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
        txt(sl, 12.25, y + 0.07, 0.45, 0.5, str(k + 1), 16, True, ACC, font=F_M, align="r")
    footer_strip(sl, 2, N)

    # 3 السعر والشمول
    sl = slide(prs, "FFFFFF"); title_bar(sl, "Price & what is included", "السعر وما يشمله")
    box(sl, 0.6, 1.7, 4.4, 4.9, NAVY)
    txt(sl, 0.8, 1.9, 4.0, 0.3, "BP-FI-05 · Service fee", 11, False, ACCD, font=F_M)
    txt(sl, 0.8, 2.3, 4.0, 1.0, "SAR %s" % n(FEE), 44, True, "FFFFFF", font=F_H)
    txt(sl, 0.8, 3.35, 4.0, 0.6, "%s ريال" % n(FEE), 28, True, "FFFFFF", rtl=True)
    txt(sl, 0.8, 4.2, 4.0, 1.2, ["Before 15%% VAT (SAR %s) · total SAR %s" % (n(VAT), n(TOTAL)), ar("قبل ضريبة 15%% (%s ريال) · الإجمالي %s ريال" % (n(VAT), n(TOTAL)))], 11, False, "EEF1FB", line=1.3)
    txt(sl, 0.8, 5.5, 4.0, 0.9, ["Government fees: **Not included**", "الرسوم الحكومية: **غير شامل**"], 12, False, ACCD, line=1.3)
    # ما يشمله
    box(sl, 5.3, 1.7, 7.43, 2.45, BG, LINE)
    txt(sl, 5.45, 1.78, 3.5, 0.3, "On us · علينا", 12, True, NAVY)
    inc = [("Eligibility study & structuring", "دراسة الأهلية والهيكلة"), ("MISA investment registration", "التسجيل الاستثماري لدى وزارة الاستثمار"),
           ("Support-entity coordination", "التنسيق مع جهة الدعم"), ("Articles & commercial registration", "عقد التأسيس والسجل التجاري"),
           ("Chamber, ZATCA, GOSI, HR, national address", "الغرفة والزكاة والتأمينات والموارد البشرية والعنوان الوطني")]
    txt(sl, 5.45, 2.1, 3.6, 2.0, ["• " + e for e, _ in inc], 10.5, False, TXT, line=1.25)
    txt(sl, 9.05, 2.1, 3.55, 2.0, [a + " •" for _, a in inc], 10.5, False, TXT, rtl=True, line=1.25)
    box(sl, 5.3, 4.3, 7.43, 2.3, "FFFFFF", LINE)
    txt(sl, 5.45, 4.38, 3.5, 0.3, "On the Client · على العميل", 12, True, NAVY)
    exc = [("Government fees & entity charges", "الرسوم الحكومية ورسوم الجهات"), ("Attestation, certified translation, support-entity fees", "التصديق والترجمة المعتمدة ورسوم جهة الدعم"),
           ("Powers of attorney (issued in the Kingdom if resident)", "الوكالات (تصدر من داخل السعودية إن كان مقيماً)"), ("Documents & information, complete and on time", "المستندات والمعلومات كاملة وفي وقتها")]
    txt(sl, 5.45, 4.7, 3.6, 1.9, ["• " + e for e, _ in exc], 10.5, False, TXT, line=1.25)
    txt(sl, 9.05, 4.7, 3.55, 1.9, [a + " •" for _, a in exc], 10.5, False, TXT, rtl=True, line=1.25)
    footer_strip(sl, 3, N)

    # 4 المراحل والمدة والدفع
    sl = slide(prs, "FFFFFF"); title_bar(sl, "Stages, duration & payment", "المراحل والمدة والدفع")
    steps = [("1", "Before registration", "قبل التسجيل", "Documents, support letter, attestation — on the Client", "المستندات وخطاب الدعم والتصديق — على العميل"),
             ("2", "Investment registration", "التسجيل الاستثماري", "Official period: 10 working days from a complete file", "المدة الرسمية: 10 أيام عمل من اكتمال الملف"),
             ("3", "Incorporation & handover", "التأسيس والتسليم", "CR, chamber, ZATCA, GOSI, HR, national address, final file", "السجل والغرفة والزكاة والتأمينات والملف النهائي")]
    for k, (no, te, ta, de, da) in enumerate(steps):
        x = 0.6 + k * 4.1
        box(sl, x, 1.7, 3.9, 2.2, BG, LINE)
        txt(sl, x + 0.15, 1.78, 0.6, 0.5, no, 22, True, ACC, font=F_M)
        txt(sl, x + 0.15, 2.3, 3.6, 0.35, te, 12.5, True, NAVY)
        txt(sl, x + 0.15, 2.65, 3.6, 0.6, de, 9.5, False, TXT, line=1.1)
        txt(sl, x + 0.15, 3.25, 3.6, 0.35, ta, 12.5, True, NAVY, rtl=True)
        txt(sl, x + 0.15, 3.58, 3.6, 0.3, ar(da), 9.5, False, TXT, rtl=True, line=1.1)
    txt(sl, 0.6, 4.15, 6, 0.4, "Payment · الدفع", 14, True, NAVY)
    for k, (pe, pa, de, da) in enumerate((("50%", "50%", "On signing the contract", "عند التعاقد"), ("50%", "50%", "After the company is established and the work is delivered", "بعد تأسيس الشركة وتسليم العمل"))):
        x = 0.6 + k * 6.15
        box(sl, x, 4.65, 5.98, 1.95, NAVY)
        txt(sl, x + 0.2, 4.75, 1.8, 0.8, pe, 34, True, "FFFFFF", font=F_H)
        txt(sl, x + 0.2, 5.55, 5.6, 0.3, "SAR %s before VAT · %s ريال قبل الضريبة" % (n(HALF), n(HALF)), 11, False, ACCD)
        txt(sl, x + 0.2, 5.9, 5.6, 0.3, de, 11, False, "FFFFFF")
        txt(sl, x + 0.2, 6.22, 5.6, 0.3, da, 12, False, "FFFFFF", rtl=True)
    footer_strip(sl, 4, N)

    # 5 الإغلاق
    sl = slide(prs, NAVY)
    box(sl, 0, 0, 13.333, 1.55, "FDFDFC")
    sl.shapes.add_picture(LOGO, Inches(0.6), Inches(0.42), width=Inches(4.2))
    txt(sl, 0.6, 2.2, 12.1, 0.3, "CONTACT & BANK TRANSFER · التواصل والتحويل البنكي", 12, False, ACCD, font=F_M)
    txt(sl, 0.6, 2.6, 12.1, 0.9, "Thank you for your trust · شكراً لثقتكم", 34, True, "FFFFFF", font=F_H)
    box(sl, 0.6, 4.0, 5.9, 0.03, ACCD); box(sl, 6.85, 4.0, 5.9, 0.03, ACCD)
    txt(sl, 0.6, 4.1, 5.9, 2.4, ["**Contact · التواصل**", "Mobile / الجوال: %s" % PHONE, "Email / الإيميل: business@businesspartner.sa", "Website / الموقع: www.businesspartner.sa",
                                "National address / العنوان الوطني: RDAA9214 · Ad Diriyah 13714"], 12, False, "FFFFFF", line=1.5)
    txt(sl, 6.85, 4.1, 5.9, 2.4, ["**Bank transfer · التحويل البنكي**", "Beneficiary / المستفيد: Business Partner Solutions / بزنس بارتنر سلوشنز", "Bank / البنك: %s" % BANK, "IBAN: %s" % IBAN], 12, False, "FFFFFF", line=1.5)
    footer_strip(sl, 5, N, dark=True)
    prs.save(out)

if __name__ == "__main__":
    from build_html import build_contract, build_nda  # noqa — لاستعمال نفس بنود الصفحات
    out = sys.argv[1]; os.makedirs(out, exist_ok=True)
    build_docx("agreement", contract_clauses(), C_TITLE_EN, C_TITLE_AR, cover_of("agreement"), os.path.join(out, "02-Agreement.docx"))
    build_docx("nda", nda_rows(), N_TITLE_EN, N_TITLE_AR, cover_of("nda"), os.path.join(out, "03-NDA.docx"))
    build_pptx(os.path.join(out, "01-Quotation-slides.pptx"))
    print("ok", out, "FINAL" if FINAL else "scrubbed")
